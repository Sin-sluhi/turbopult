# -*- coding: utf-8 -*-
"""Сборка базы объяснений ошибок checkxml.gge.ru для пульта.

Берёт только проверенный материал из базы знаний ИИ-помощника
(C:\\TelegramBot\\AI_KB) — ничего не придумывает:
  * рецепты «Ошибка_NN_*.txt» — тексты ошибок, шаги в Турбосметчике;
  * «Проверки_КПСР\\Ошибки_КПСР_и_решения.xlsx» — что значит проверка
    сверки с базой и как её устранить;
  * XSD-схемы ГГЭ («Разбор_EvhEstGGE\\xsd») — русское описание каждого поля.
Результат — backend/app/gge_kb.json. Запуск: python backend\\build_gge_kb.py
"""
import json
import pathlib
import re
import sys
import xml.etree.ElementTree as ET

KB = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else r"C:\TelegramBot\AI_KB")
OUT = pathlib.Path(__file__).resolve().parent / "app" / "gge_kb.json"
XS = "{http://www.w3.org/2001/XMLSchema}"


def recipes():
    out = []
    for f in sorted((KB / "Инструкции_со_скриншотами").glob("Ошибка_*.txt")):
        num = int(re.match(r"Ошибка_(\d+)", f.name).group(1))
        if not (1 <= num <= 17 or 50 <= num <= 90):
            continue          # 18–49 — сбои программы, не про XML
        text = f.read_text(encoding="utf-8")
        lines = text.splitlines()
        title = lines[0].lstrip("# ").strip()
        steps = [re.sub(r"\s+", " ", m).strip()
                 for m in re.findall(r"^Шаг \d+\.\s*(.+?)(?=^\[ФОТО|^Шаг \d+\.|^Важно:|^Коротко:|\Z)", text, re.S | re.M)]
        important = re.search(r"^Важно:\s*(.+?)(?=^Коротко:|\Z)", text, re.S | re.M)
        short = re.search(r"^Коротко:\s*(.+)\Z", text, re.S | re.M)
        short_items = [re.sub(r"^\d+\.\s*", "", s).strip()
                       for s in (short.group(1).splitlines() if short else []) if re.match(r"^\d+\.", s.strip())]
        # тексты ошибок и узлы: «В узле "P" отсутствует "C"», "P/C", «элемент "C"»
        quoted = re.findall(r"«([^»]+)»", title) + re.findall(r"^\s{2}(XML-файл.+)$", text, re.M)
        pairs = set()
        for q in quoted + [title]:
            for p, c in re.findall(r'[Уу]зле "([\w/\\]+)"[^"]*?"(\w+)"', q):
                parts = re.split(r"[/\\]", p)
                pairs.add((parts[-1], c))
                if len(parts) > 1:
                    pairs.add((parts[-2], parts[-1]))
            for p in re.findall(r'[Уу]зле "(\w+)[/\\](\w+)"', q):
                pairs.add(p)
            for c in re.findall(r'элемент "(\w+)"', q):
                pairs.add(("", c))
            for c in re.findall(r'узла? "(\w+)"', q):
                pairs.add((c, ""))
        out.append({"num": num, "file": f.name, "title": title,
                    "texts": quoted, "pairs": sorted(pairs),
                    "steps": steps,
                    "important": re.sub(r"\s+", " ", important.group(1)).strip() if important else "",
                    "short": short_items})
    return out


def kpsr():
    try:
        import openpyxl
    except ImportError:
        print("нет openpyxl — таблица КПСР пропущена (pip install openpyxl)")
        return [], []
    wb = openpyxl.load_workbook(KB / "Проверки_КПСР" / "Ошибки_КПСР_и_решения.xlsx", read_only=True)
    rows = list(wb["Ошибки КПСР"].iter_rows(values_only=True))[1:]
    checks = [{"key": r[0], "stage": r[1], "group": r[2], "name": r[3], "verbatim": r[4] or "",
               "severity": r[5] or "", "meaning": r[6] or "", "solution": r[7] or ""}
              for r in rows if r and r[0]]
    groups = [{"stage": r[0], "group": r[1], "what": r[2], "how": r[3]}
              for r in list(wb["Группы проверок"].iter_rows(values_only=True))[1:] if r and r[1]]
    return checks, groups


def xsd_docs():
    """Имя элемента -> [{path, doc}] по всем схемам."""
    docs = {}

    def walk(node, path):
        for el in node.findall(f"{XS}element"):
            name = el.get("name")
            if not name:
                continue
            p = path + [name]
            doc = el.find(f"{XS}annotation/{XS}documentation")
            if doc is not None and doc.text and doc.text.strip():
                text = re.sub(r"\s+", " ", doc.text).strip().strip('"')
                lst = docs.setdefault(name, [])
                key = "/".join(p[-3:])
                if not any(d["path"] == key for d in lst):
                    lst.append({"path": key, "doc": text})
            walk(el, p)
        for tag in ("complexType", "sequence", "choice", "all"):
            for sub in node.findall(f"{XS}{tag}"):
                walk(sub, path)

    for f in sorted((KB / "Разбор_EvhEstGGE" / "xsd").glob("*.xsd")):
        try:
            root = ET.parse(f).getroot()
        except ET.ParseError:
            continue
        walk(root, [])
        # именованные типы: их поля тоже описаны
        for ct in root.findall(f"{XS}complexType"):
            walk(ct, [ct.get("name") or ""])
    return docs


def main():
    r = recipes()
    checks, groups = kpsr()
    d = xsd_docs()
    OUT.write_text(json.dumps({"recipes": r, "kpsr": checks, "groups": groups, "xsd": d},
                              ensure_ascii=False, indent=0), encoding="utf-8")
    print(f"рецептов {len(r)}, проверок КПСР {len(checks)}, групп {len(groups)}, "
          f"полей XSD {len(d)} -> {OUT}")


if __name__ == "__main__":
    main()
